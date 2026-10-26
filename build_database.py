import sqlite3
import openpyxl
import os

DB_FILE = "food.db"
EXCEL_FILE = "food_ingredients.xlsx"

def needs_sync():
    """Return True if the Excel file is newer than the database, or DB doesn't exist yet."""
    if not os.path.exists(DB_FILE):
        return True
    excel_time = os.path.getmtime(EXCEL_FILE)
    db_time = os.path.getmtime(DB_FILE)
    return excel_time > db_time

def sync_from_excel():
    """Rebuild food.db from food_ingredients.xlsx."""
    conn = sqlite3.connect(DB_FILE)
    cur = conn.cursor()

    cur.execute("""
        CREATE TABLE IF NOT EXISTS additives (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            name TEXT NOT NULL,
            aliases TEXT,
            category TEXT NOT NULL,
            reason TEXT
        )
    """)

    cur.execute("""
        CREATE TABLE IF NOT EXISTS nutrient_thresholds (
            nutrient TEXT PRIMARY KEY,
            low_max REAL,
            medium_max REAL,
            unit TEXT,
            note TEXT
        )
    """)

    cur.execute("DELETE FROM additives")
    cur.execute("DELETE FROM nutrient_thresholds")

    wb = openpyxl.load_workbook(EXCEL_FILE, data_only=True)

    ws = wb["additives"]
    count_additives = 0
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] is None:
            continue
        name, aliases, category, reason = row
        cur.execute(
            "INSERT INTO additives (name, aliases, category, reason) VALUES (?, ?, ?, ?)",
            (name, aliases or "", category, reason)
        )
        count_additives += 1

    ws2 = wb["nutrient_thresholds"]
    count_thresholds = 0
    for row in ws2.iter_rows(min_row=2, values_only=True):
        if row[0] is None:
            continue
        nutrient, low_max, medium_max, unit, note = row
        cur.execute(
            "INSERT INTO nutrient_thresholds (nutrient, low_max, medium_max, unit, note) VALUES (?, ?, ?, ?, ?)",
            (nutrient, low_max, medium_max, unit, note)
        )
        count_thresholds += 1

    conn.commit()
    conn.close()

    print(f"Synced: {count_additives} additives, {count_thresholds} nutrient thresholds -> {DB_FILE}")

def ensure_synced():
    """Call this from app.py on startup — only rebuilds if Excel changed."""
    if needs_sync():
        sync_from_excel()
    else:
        print("Database already up to date, skipping import.")

if __name__ == "__main__":
    sync_from_excel()